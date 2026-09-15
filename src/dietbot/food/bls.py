"""Import tabeli BLS 4.0 (Max Rubner-Institut, CC BY 4.0) do lekkiego indeksu JSON.

Surowy plik XLSX (14 MB, 7140 produktów × 418 kolumn) leży w dane/raw/bls4 (poza gitem).
Indeks zawiera tylko potrzebne składniki, nazwy DE/EN i pochodzenie każdej wartości.
"""

from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path

BLS_VERSION = "4.0 (2025)"
BLS_LICENSE = "CC BY 4.0 — Max Rubner-Institut, Bundeslebensmittelschlüssel (BLS) 4.0"

# kod składnika BLS -> pole w naszym modelu
COMPONENTS: dict[str, str] = {
    "ENERCC": "kcal",
    "PROT625": "protein",
    "FAT": "fat",
    "CHO": "carbs",
    "FIBT": "fiber",
    "SUGAR": "sugars",
    "FASAT": "saturated_fat",
    "NACL": "salt",
    "ALC": "alcohol",
    "OA": "organic_acids",
    "POLYL": "polyols",
    "WATER": "water",
}

GROUPS_PL: dict[str, str] = {
    "B": "pieczywo",
    "C": "zboża, mąki, ryż",
    "D": "ciasta i wypieki",
    "E": "makarony, kluski, jaja",
    "F": "owoce",
    "G": "warzywa",
    "H": "strączki, orzechy, nasiona, zamienniki roślinne",
    "K": "ziemniaki i grzyby",
    "M": "mleko i nabiał, sery",
    "N": "napoje bezalkoholowe",
    "P": "napoje alkoholowe",
    "Q": "tłuszcze i oleje",
    "R": "przyprawy, sosy, dodatki do wypieków",
    "S": "cukier, słodycze, lody, miód",
    "T": "ryby i owoce morza",
    "U": "mięso",
    "V": "drób, dziczyzna, podroby",
    "W": "wędliny i kiełbasy",
    "X": "potrawy (receptury)",
    "Y": "potrawy (receptury)",
}

# Oznaczenia śladowych ilości -> 0; brak danych -> None
_TRACE = re.compile(r"^(tr|<lod|<loq|<lod or <loq|spuren)$", re.IGNORECASE)


def parse_value(raw: object) -> tuple[float | None, str | None]:
    """Zwraca (wartość, flaga). Flaga: 'slad' dla śladów, 'brak' dla braku danych."""
    if raw is None:
        return None, "brak"
    if isinstance(raw, int | float):
        return float(raw), None
    text = str(raw).strip()
    if text in {"", "-"}:
        return None, "brak"
    if _TRACE.match(text):
        return 0.0, "slad"
    try:
        return float(text.replace(",", ".")), None
    except ValueError:
        return None, "brak"


@dataclass
class BlsFood:
    code: str
    name_de: str
    name_en: str
    values: dict[str, float | None]
    origins: dict[str, str | None] = field(default_factory=dict)  # Datenherkunft, np. Analyse, Literatur
    flags: dict[str, str] = field(default_factory=dict)

    @property
    def group(self) -> str:
        return self.code[0]

    @property
    def is_dish(self) -> bool:
        return self.group in ("X", "Y")


def import_xlsx(xlsx_path: Path) -> list[BlsFood]:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    value_col: dict[str, int] = {}
    origin_col: dict[str, int] = {}
    for i, h in enumerate(header):
        if not h:
            continue
        code = str(h).split(" ")[0]
        if code not in COMPONENTS:
            continue
        if str(h).endswith("Datenherkunft"):
            origin_col[code] = i
        elif not str(h).endswith("Referenz"):
            value_col[code] = i
    missing = set(COMPONENTS) - set(value_col)
    if missing:
        raise ValueError(f"w pliku BLS brakuje kolumn: {sorted(missing)}")
    foods: list[BlsFood] = []
    for r in rows:
        if not r or not r[0]:
            continue
        values: dict[str, float | None] = {}
        origins: dict[str, str | None] = {}
        flags: dict[str, str] = {}
        for code, field_name in COMPONENTS.items():
            v, flag = parse_value(r[value_col[code]])
            values[field_name] = v
            if flag:
                flags[field_name] = flag
            if code in origin_col:
                origins[field_name] = r[origin_col[code]]
        foods.append(BlsFood(code=str(r[0]), name_de=str(r[1]), name_en=str(r[2]), values=values,
                             origins=origins, flags=flags))
    return foods


def write_index(foods: list[BlsFood], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"version": BLS_VERSION, "license": BLS_LICENSE, "foods": [asdict(f) for f in foods]}
    out_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def _fold(text: str) -> str:
    text = unicodedata.normalize("NFKD", text.lower())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text.replace("ß", "ss")


class BlsIndex:
    def __init__(self, foods: list[BlsFood], version: str = BLS_VERSION):
        self.foods = foods
        self.version = version
        self.by_code = {f.code: f for f in foods}
        self._folded = [(_fold(f.name_de), _fold(f.name_en)) for f in foods]

    @classmethod
    def load(cls, index_path: Path) -> BlsIndex:
        payload = json.loads(index_path.read_text(encoding="utf-8"))
        foods = [BlsFood(**f) for f in payload["foods"]]
        return cls(foods, payload.get("version", BLS_VERSION))

    def get(self, code: str) -> BlsFood:
        try:
            return self.by_code[code]
        except KeyError:
            raise KeyError(f"brak kodu BLS {code}") from None

    def search(self, keywords: list[str], groups: str | None = None, exclude: list[str] | None = None,
               include_dishes: bool = False, limit: int = 15) -> list[BlsFood]:
        """Wszystkie słowa kluczowe (DE lub EN) muszą wystąpić w nazwie. Bez potraw X/Y, chyba że wskazano."""
        kws = [_fold(k) for k in keywords if k.strip()]
        exc = [_fold(k) for k in (exclude or []) if k.strip()]
        out: list[BlsFood] = []
        for food, (de, en) in zip(self.foods, self._folded, strict=True):
            if groups and food.group not in groups:
                continue
            if not include_dishes and not groups and food.is_dish:
                continue
            text = f"{de} | {en}"
            if all(k in text for k in kws) and not any(e in text for e in exc):
                out.append(food)
        # najpierw krótsze nazwy (zwykle produkt podstawowy, a nie wariant)
        out.sort(key=lambda f: (len(f.name_de), f.code))
        return out[:limit]
